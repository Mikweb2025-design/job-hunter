"""CSV / XLSX export of the tracker."""
from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("id", "ID"), ("score", "Score"), ("rule_score", "Regel-Score"), ("llm_score", "LLM-Score"),
    ("status", "Status"), ("applied_date", "Beworben am"), ("title", "Titel"),
    ("company", "Unternehmen"), ("location", "Ort"), ("remote", "Remote"),
    ("salary_min", "Gehalt min"), ("salary_max", "Gehalt max"), ("source", "Quelle"),
    ("published", "Veröffentlicht"), ("fetched_at", "Gefunden"), ("url", "Link"),
    ("llm_reason", "Begründung"), ("notes", "Notizen"),
    ("channel_label", "Kanal"), ("interview_at", "Gespräch am"), ("follow_up_at", "Nachfassen am"),
    ("close_reason", "Abschlussgrund"),
]
STATUS_LABELS = {"neu": "Neu", "interessant": "Interessant", "beworben": "Beworben", "gespraech": "Gespräch",
                 "absage": "Absage", "angebot": "Angebot", "zu_weit": "Zu weit"}


def _safe(value):
    # Postings are third-party text: neutralise spreadsheet formulas (CSV injection).
    if isinstance(value, str) and value[:1] in ("=", "+", "@", "\t", "\r"):
        return "'" + value
    return value


def _value(j: dict, key: str):
    if key == "status":
        return STATUS_LABELS.get(j.get("status") or "", j.get("status"))
    return j.get(key)


def _rows(jobs: list[dict]):
    for j in jobs:
        yield [_safe(_value(j, k)) if _value(j, k) is not None else "" for k, _ in COLUMNS]


def to_csv(jobs: list[dict]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([label for _, label in COLUMNS])
    w.writerows(_rows(jobs))
    return "﻿" + buf.getvalue()  # BOM so Excel opens UTF-8 correctly


def to_xlsx(jobs: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bewerbungen"
    ws.append([label for _, label in COLUMNS])
    for row in _rows(jobs):
        ws.append(row)
    widths = {"Kanal": 30, "Titel": 45, "Unternehmen": 28, "Link": 40, "Begründung": 60, "Notizen": 40}
    for i, (_, label) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(label, 14)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
